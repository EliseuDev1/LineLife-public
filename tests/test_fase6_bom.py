"""Testes automatizados da Fase 6: Lista de Materiais Inteligente (Bill of Materials - BOM).

Valida:
1. Quantificação exata de cabo com margens de segurança e terminações.
2. Dimensionamento de postes terminais e intermediários.
3. Dimensionamento de esticadores (simples vs duplo para linhas longas / múltiplos vãos).
4. Escalonamento de grampos pesados (clips) em função do diâmetro do cabo (NBR ISO 2408 / US Fed Spec).
5. Variação dos elementos de fixação por tipo de instalação (Laje, Metálica, Telhado, Munck).
6. Trava de segurança (BOM só gera se projeto aprovado e sem pendências).
7. Geração e integridade da planilha Excel (.xlsx) formatada.
"""

import os
import tempfile
import unittest
from unittest.mock import Mock
import openpyxl

from src.application.project_service import ProjectService
from src.catalogs.cable_catalog import CableCatalog
from src.core.license_models import LicenseData, LicenseTier
from src.catalogs.material_rules import BOMGenerator, BOMItem, BillOfMaterials
from src.domain.installation_types import (
    IndustrialInstallation,
    MunckInstallation,
    RoofInstallation,
    SlabInstallation,
)
from src.domain.models import (
    CableProperties,
    FallClearanceInput,
    LoadRuleType,
    SpanGeometry,
    UserLoadConfig,
)
from src.domain.project import Project, ProjectStatus


class TestFase6BOM(unittest.TestCase):
    """Bateria de testes unitários para a Lista de Materiais (BOM)."""

    def setUp(self):
        self.catalog = CableCatalog()
        license_manager = Mock(get_license=lambda: LicenseData(tier=LicenseTier.LOCAL))
        self.service = ProjectService(self.catalog, license_manager=license_manager)
        self.cabo_8mm = self.catalog.get_by_specification("Ø5/16 - 6X19 SEALE EIPS - AF")
        self.cabo_10mm = self.catalog.get_by_specification("Ø3/8 - 6X19 SEALE EIPS - AF")
        self.cabo_13mm = self.catalog.get_by_specification("Ø1/2 - 6X19 SEALE EIPS - AF")

    def _create_approved_project(
        self,
        num_spans: int = 1,
        span_length_mm: float = 6600.0,
        cable: CableProperties = None,
        installation=None,
    ) -> Project:
        """Cria e calcula um projeto válido e aprovado para testes da BOM."""
        if cable is None:
            cable = self.cabo_8mm
        if installation is None:
            installation = SlabInstallation(
                structure_description="Viga de Concreto 30x50",
                technical_notes="Concreto fck 30 MPa",
                is_structure_validated_by_engineer=True,
                responsible_engineer_art="ART-TESTE-123456",
                slab_type="Concreto Armado Maciço",
                thickness_mm=150.0,
                concrete_fck_mpa=30.0,
                anchor_type="Chumbador Químico M12",
                anchor_edge_distance_mm=150.0,
            )

        spans = [
            SpanGeometry(span_index=i + 1, span_mm=span_length_mm, initial_sag_percent=3.0)
            for i in range(num_spans)
        ]

        proj = self.service.create_project(
            title="Projeto Linha de Teste BOM",
            client="Cliente Teste",
            location="Galpão 01",
            responsible_engineer="Eng. Teste",
            installation=installation,
        )
        if cable:
            proj.cable = cable
        proj.user_config = UserLoadConfig(total_users=1, simultaneous_users=1)
        proj.clearance_input = FallClearanceInput(available_clearance_m=10.0)
        proj.update_spans(spans)

        self.service.calculate_project(proj)
        return proj

    def test_bom_generation_single_span(self):
        """Valida a lista de materiais para um vão único curto (6.6 m)."""
        proj = self._create_approved_project(num_spans=1, span_length_mm=6600.0)
        self.assertTrue(proj.can_generate_bom)

        bom = proj.generate_bom()
        self.assertEqual(bom.project_title, proj.title)
        self.assertEqual(bom.total_posts_count, 2)
        self.assertAlmostEqual(bom.total_spans_length_m, 6.6, places=2)
        # 6.6 + 4.0 = 10.6 m -> arredondado para múltiplo de 0.5 = 11.0 m
        self.assertEqual(bom.cable_total_length_m, 11.0)

        codes = {it.item_code: it for it in bom.items}
        self.assertIn("CAB-001", codes)
        self.assertEqual(codes["CAB-001"].quantity, 11.0)
        self.assertEqual(codes["CAB-001"].unit, "m")

        self.assertIn("PST-EXT-01", codes)
        self.assertEqual(codes["PST-EXT-01"].quantity, 2.0)

        # Sem postes intermediários para 1 vão
        self.assertNotIn("PST-INT-01", codes)

        # Esticador: 1 para vão curto
        self.assertIn("EST-001", codes)
        self.assertEqual(codes["EST-001"].quantity, 1.0)

        # Sapatilhos: 2
        self.assertIn("SAP-001", codes)
        self.assertEqual(codes["SAP-001"].quantity, 2.0)

        # Grampos para 8.0 mm: 3 por ponta = 6
        self.assertIn("GRP-001", codes)
        self.assertEqual(codes["GRP-001"].quantity, 6.0)

        # Absorvedor: 1
        self.assertIn("ABS-001", codes)
        self.assertEqual(codes["ABS-001"].quantity, 1.0)

        # Fixação Laje: 2 postes * 4 chumbadores = 8
        self.assertIn("FIX-SLAB-01", codes)
        self.assertEqual(codes["FIX-SLAB-01"].quantity, 8.0)

        # Placa NR-35: 1
        self.assertIn("PLC-001", codes)
        self.assertEqual(codes["PLC-001"].quantity, 1.0)

    def test_bom_generation_multi_span_long(self):
        """Valida linha longa com múltiplos vãos (4 vãos de 10 m = 40 m, 5 postes)."""
        proj = self._create_approved_project(num_spans=4, span_length_mm=10000.0)
        bom = proj.generate_bom()

        self.assertEqual(bom.total_posts_count, 5)
        self.assertAlmostEqual(bom.total_spans_length_m, 40.0, places=2)
        # 40.0 + 4.0 = 44.0 m
        self.assertEqual(bom.cable_total_length_m, 44.0)

        codes = {it.item_code: it for it in bom.items}
        # Postes terminais = 2
        self.assertEqual(codes["PST-EXT-01"].quantity, 2.0)
        # Postes intermediários = 3
        self.assertIn("PST-INT-01", codes)
        self.assertEqual(codes["PST-INT-01"].quantity, 3.0)

        # Linha > 30 m e >= 4 vãos exige 2 esticadores
        self.assertEqual(codes["EST-001"].quantity, 2.0)

        # Fixação Laje para 5 postes * 4 = 20 chumbadores
        self.assertEqual(codes["FIX-SLAB-01"].quantity, 20.0)

    def test_bom_clip_scaling_by_diameter(self):
        """Valida o escalonamento do número de grampos por diâmetro de cabo."""
        # 1. Cabo 8.0 mm (<= 8.0 mm) -> 3 por ponta = 6
        p1 = self._create_approved_project(cable=self.cabo_8mm)
        b1 = p1.generate_bom()
        c1 = {it.item_code: it for it in b1.items}
        self.assertEqual(c1["GRP-001"].quantity, 6.0)

        # 2. Cabo 9.5 mm (<= 12.0 mm) -> 4 por ponta = 8
        p2 = self._create_approved_project(cable=self.cabo_10mm)
        b2 = p2.generate_bom()
        c2 = {it.item_code: it for it in b2.items}
        self.assertEqual(c2["GRP-001"].quantity, 8.0)

        # 3. Cabo 13.0 mm (> 12.0 mm) -> 5 por ponta = 10
        p3 = self._create_approved_project(cable=self.cabo_13mm)
        b3 = p3.generate_bom()
        c3 = {it.item_code: it for it in b3.items}
        self.assertEqual(c3["GRP-001"].quantity, 10.0)

    def test_bom_fasteners_by_installation_category(self):
        """Valida que cada categoria de instalação seleciona os fixadores corretos."""
        # Industrial (Metálica): Beam Clamps
        inst_ind = IndustrialInstallation(
            structure_description="Viga I W250x32",
            is_structure_validated_by_engineer=True,
            responsible_engineer_art="ART-MET-01",
        )
        p_ind = self._create_approved_project(num_spans=2, span_length_mm=5000.0, installation=inst_ind)
        b_ind = p_ind.generate_bom()
        codes_ind = {it.item_code: it for it in b_ind.items}
        self.assertIn("FIX-IND-01", codes_ind)
        # 3 postes * 4 grampos de aba = 12
        self.assertEqual(codes_ind["FIX-IND-01"].quantity, 12.0)

        # Roof (Telhado): Kit Interface Cobertura
        inst_roof = RoofInstallation(
            structure_description="Terça U 150x50",
            support_spacing_mm=1200.0,
            is_structure_validated_by_engineer=True,
            responsible_engineer_art="ART-ROOF-01",
        )
        p_roof = self._create_approved_project(num_spans=2, span_length_mm=5000.0, installation=inst_roof)
        b_roof = p_roof.generate_bom()
        codes_roof = {it.item_code: it for it in b_roof.items}
        self.assertIn("FIX-ROOF-01", codes_roof)
        # 3 postes = 3 kits
        self.assertEqual(codes_roof["FIX-ROOF-01"].quantity, 3.0)

        # Munck: Kit Mancal Articulado Munck
        inst_mnc = MunckInstallation(
            structure_description="Braço Guindauto",
            is_structure_validated_by_engineer=True,
            responsible_engineer_art="ART-MNC-01",
        )
        p_mnc = self._create_approved_project(num_spans=1, span_length_mm=5000.0, installation=inst_mnc)
        b_mnc = p_mnc.generate_bom()
        codes_mnc = {it.item_code: it for it in b_mnc.items}
        self.assertIn("FIX-MNC-01", codes_mnc)
        self.assertEqual(codes_mnc["FIX-MNC-01"].quantity, 2.0)

    def test_bom_security_lock(self):
        """Valida que a BOM recusa geração caso o projeto não esteja apto."""
        proj = self._create_approved_project()

        # 1. Status DRAFT
        proj.status = ProjectStatus.DRAFT
        self.assertFalse(proj.can_generate_bom)
        with self.assertRaises(ValueError):
            proj.generate_bom()

        # 2. Status CALCULATED_REJECTED
        proj.status = ProjectStatus.CALCULATED_REJECTED
        self.assertFalse(proj.can_generate_bom)
        with self.assertRaises(ValueError):
            proj.generate_bom()

        # 3. Status INVALIDATED
        proj.status = ProjectStatus.INVALIDATED
        self.assertFalse(proj.can_generate_bom)
        with self.assertRaises(ValueError):
            proj.generate_bom()

        # 4. Falta de ART / Validação Estrutural
        proj.status = ProjectStatus.CALCULATED_VALID
        proj.installation.is_structure_validated_by_engineer = False
        proj.installation.responsible_engineer_art = None
        self.assertFalse(proj.can_generate_bom)
        with self.assertRaises(ValueError):
            proj.generate_bom()

    def test_bom_excel_export_and_content(self):
        """Valida a exportação completa da BOM para arquivo Excel (.xlsx)."""
        proj = self._create_approved_project(num_spans=3, span_length_mm=8000.0)
        bom = proj.generate_bom()

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            out_file = BOMGenerator.export_to_excel(bom, tmp_path)
            self.assertTrue(os.path.exists(out_file))
            self.assertGreater(os.path.getsize(out_file), 1000)

            # Abre com openpyxl e valida o conteúdo gravado
            wb = openpyxl.load_workbook(out_file)
            self.assertIn("Lista de Materiais - BOM", wb.sheetnames)
            ws = wb["Lista de Materiais - BOM"]

            # Título na A1
            self.assertIn("LISTA DE MATERIAIS TÉCNICA", ws["A1"].value)

            # Cabeçalhos na linha 8
            self.assertEqual(ws["A8"].value, "Item")
            self.assertEqual(ws["B8"].value, "Código")
            self.assertEqual(ws["C8"].value, "Categoria")
            self.assertEqual(ws["E8"].value, "Qtd.")
            self.assertEqual(ws["F8"].value, "Un.")

            # Verifica se os itens estão presentes a partir da linha 9
            row_count = 0
            for r in range(9, 9 + len(bom.items)):
                self.assertIsNotNone(ws[f"B{r}"].value)
                row_count += 1
            self.assertEqual(row_count, len(bom.items))

            wb.close()
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)


if __name__ == "__main__":
    unittest.main()
